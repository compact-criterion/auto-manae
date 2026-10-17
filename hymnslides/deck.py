"""Build the PowerPoint deck: a title slide, then music slides in singing order."""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .engrave import MIN_LYRIC_PT, Engraving, Line, engrave, lines_for
from .songformat import Song, SongError, check_copyright

SLIDE_W, SLIDE_H = 13.333, 7.5
FONT = "Segoe UI"
LABEL_PT = 30
# Where the music goes, below the section label (left, top, width, height in inches).
MUSIC_AREA = (0.2, 0.95, 12.93, 6.4)
PANEL_PAD = 0.3


@dataclass
class Theme:
    name: str
    background: str
    music: str
    text: str
    panel: bool


def make_theme(name: str, bg: str | None = None) -> Theme:
    if name == "dark":
        return Theme("dark", "#1C1C1C", "#FFFFFF", "#FFFFFF", False)
    if name == "light":
        return Theme("light", "#FFFFFF", "#111111", "#111111", False)
    if name == "color":
        return Theme("color", bg or "#8B1A1A", "#111111", "#FFFFFF", True)
    raise SongError(f'Unknown theme "{name}". Use dark, light or color.')


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _music_box(theme: Theme) -> tuple[float, float, float, float]:
    left, top, w, h = MUSIC_AREA
    if theme.panel:
        return left + PANEL_PAD, top + PANEL_PAD, w - 2 * PANEL_PAD, h - 2 * PANEL_PAD
    return MUSIC_AREA


# ---------------------------------------------------------------- slide parts

def _new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def _blank_slide(prs: Presentation, theme: Theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(theme.background)
    return slide


def _text_box(slide, left, top, width, height, paragraphs, color, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP):
    """paragraphs: list of (text, size_pt, bold, italic)."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, (text, size, bold, italic) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(size * 0.35)
        run = p.add_run()
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = _rgb(color)
    return box


def _panel(slide):
    left, top, w, h = MUSIC_AREA
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                   Inches(w), Inches(h))
    shape.adjustments[0] = 0.05
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _label(slide, text: str, theme: Theme):
    _text_box(slide, 0.35, 0.12, 9.0, 0.75, [(text, LABEL_PT, True, False)], theme.text,
              anchor=MSO_ANCHOR.MIDDLE)


def title_slide(prs, theme: Theme, meta: dict) -> None:
    slide = _blank_slide(prs, theme)
    paras = []
    if str(meta.get("number") or "").strip():
        paras.append((f"Hymn {meta['number']}", 32, False, False))
    paras.append((str(meta.get("title", "")), 54, True, False))
    if meta.get("scripture"):
        paras.append((str(meta["scripture"]), 32, False, True))
    if meta.get("words"):
        paras.append((f"Words: {meta['words']}", 30, False, False))
    if meta.get("music"):
        paras.append((f"Music: {meta['music']}", 30, False, False))
    if meta.get("copyright"):
        paras.append((str(meta["copyright"]), 30, False, False))
    ccli = []
    if meta.get("ccli_song"):
        ccli.append(f"CCLI Song # {meta['ccli_song']}")
    if meta.get("ccli_license"):
        ccli.append(f"CCLI License # {meta['ccli_license']}")
    if ccli:
        paras.append(("   ".join(ccli), 30, False, False))
    _text_box(slide, 0.6, 0.4, SLIDE_W - 1.2, SLIDE_H - 0.8, paras, theme.text,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def _plain_words(text: str | None) -> str:
    if not text:
        return ""
    return " ".join(w.replace("-", "") for w in text.split() if w != "_")


# ---------------------------------------------------------------- music slides

@dataclass
class MusicSlide:
    label: str
    lines: list[Line]
    engraving: Engraving
    final: bool
    lyric_pt: float = 0.0


def _engrave_checked(song, lines, theme, box, final) -> tuple[Engraving, float]:
    result = engrave(song, lines, theme.music, box, final)
    return result, result.lyric_pt(box)


def plan_section(song: Song, index: int, theme: Theme) -> list[MusicSlide]:
    """Two systems per slide; fewer measures per system if lyrics would be under 30 pt."""
    sec = song.structure[index]
    box = _music_box(theme)[2:]
    lines = lines_for(song, index)
    slides: list[MusicSlide] = []
    pairs = [lines[i:i + 2] for i in range(0, len(lines), 2)]
    for pi, pair in enumerate(pairs):
        final = pi == len(pairs) - 1
        result, pt = _engrave_checked(song, pair, theme, box, final)
        if pt >= MIN_LYRIC_PT:
            slides.append(MusicSlide(sec.label, pair, result, final, pt))
            continue
        # Too small: one phrase per slide, split in two systems when needed.
        for li, line in enumerate(pair):
            line_final = final and li == len(pair) - 1
            result, pt = _engrave_checked(song, [line], theme, box, line_final)
            if pt >= MIN_LYRIC_PT:
                slides.append(MusicSlide(sec.label, [line], result, line_final, pt))
                continue
            halves = line.split(song)
            split, split_pt = _engrave_checked(song, halves, theme, box, line_final)
            if split_pt >= MIN_LYRIC_PT or len(halves) == 1:
                slides.append(MusicSlide(sec.label, [line], split, line_final, split_pt))
                continue
            # Still too small: each half of the line on its own slide.
            for hi, half in enumerate(halves):
                half_final = line_final and hi == len(halves) - 1
                result, pt = _engrave_checked(song, [half], theme, box, half_final)
                slides.append(MusicSlide(sec.label, [half], result, half_final, pt))
    return slides


def _music_slide(prs, theme: Theme, ms: MusicSlide):
    slide = _blank_slide(prs, theme)
    _label(slide, ms.label, theme)
    if theme.panel:
        _panel(slide)
    left, top, bw, bh = _music_box(theme)
    png = ms.engraving.png
    w_in = min(bw, bh * png.width / png.height)
    h_in = w_in * png.height / png.width
    buf = io.BytesIO()
    png.save(buf, format="PNG")
    buf.seek(0)
    slide.shapes.add_picture(buf, Inches(left + (bw - w_in) / 2), Inches(top + (bh - h_in) / 2),
                             Inches(w_in), Inches(h_in))
    notes = []
    seen = set()
    for line in ms.lines:
        if (line.phrase, line.lyrics) in seen:
            continue
        seen.add((line.phrase, line.lyrics))
        notes.append(_plain_words(line.lyrics))
        if line.alto:
            notes.append("(alto) " + _plain_words(line.alto))
    slide.notes_slide.notes_text_frame.text = f"{ms.label}\n" + "\n".join(notes)
    return slide


def build_deck(song: Song, out: str | Path, theme_name: str = "light", bg: str | None = None,
               log=print) -> list[MusicSlide]:
    check_copyright(song.meta, engraving=True)
    if song.meta.get("reviewed") is False:
        raise SongError("This song was imported automatically and has not been checked yet. "
                        "Open its review.html, compare every line with the page, fix the song file, "
                        "and then change 'reviewed: false' to 'reviewed: true'.")
    theme = make_theme(theme_name, bg)
    prs = _new_deck()
    meta = dict(song.meta, title=song.title)
    title_slide(prs, theme, meta)
    planned: list[MusicSlide] = []
    cache: dict[tuple, list[MusicSlide]] = {}
    for i, sec in enumerate(song.structure):
        key = (sec.label, tuple(sec.phrases), sec.lyrics)
        if key not in cache:  # the refrain is engraved once and reused
            cache[key] = plan_section(song, i, theme)
        planned.extend(cache[key])
    for n, ms in enumerate(planned, start=2):
        _music_slide(prs, theme, ms)
        flag = "" if ms.lyric_pt >= MIN_LYRIC_PT else "  <-- under 30 pt"
        log(f"  slide {n:2d}  {ms.label:<10} lyrics {ms.lyric_pt:4.1f} pt{flag}")
    prs.save(out)
    return planned


# ---------------------------------------------------------------- template mode

def parse_sections(spec: str) -> list[tuple[str, int]]:
    """"Stanza 1:2,Refrain:2" -> [("Stanza 1", 2), ("Refrain", 2)]"""
    out = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        label, _, count = part.rpartition(":")
        if not label or not count.strip().isdigit():
            raise SongError(f'"{part}" should look like "Stanza 1:2" (label, colon, number of slides).')
        out.append((label.strip(), int(count)))
    return out


def build_template(meta: dict, sections: str, out: str | Path, theme_name: str = "color",
                   bg: str | None = None) -> int:
    check_copyright(meta, engraving=False)
    theme = make_theme(theme_name, bg)
    prs = _new_deck()
    title_slide(prs, theme, meta)
    count = 0
    for label, n in parse_sections(sections):
        for _ in range(n):
            slide = _blank_slide(prs, theme)
            _label(slide, label, theme)
            _panel(slide)
            left, top, w, h = MUSIC_AREA
            _text_box(slide, left, top, w, h, [("Paste 2 music systems here", 30, False, False)],
                      "#8A8A8A", align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            slide.notes_slide.notes_text_frame.text = label
            count += 1
    prs.save(out)
    return count
