"""Picture slides: put the printed lines of music from a scan straight onto slides.

No notes are read, so nothing can be misread: each system (treble + bass staff
with its words) is cut out of the page, cleaned up, recoloured for the theme,
and placed two to a slide. The words stay the size the hymnal printed them, so
they come out smaller than on engraved slides.
"""

from __future__ import annotations

import io
import tempfile
from pathlib import Path

from PIL import Image, ImageOps
from pptx.util import Inches

from . import omr
from .deck import _blank_slide, _label, _music_box, _new_deck, _panel, make_theme, title_slide
from .songformat import SongError, check_copyright

GAP = 40  # pixels between two systems on one slide


def systems_from_scans(sources: list[Path], pages: str | None, work: Path, log=print) -> list[Image.Image]:
    """Every system on the given pages, top to bottom, page by page, trimmed to the ink."""
    systems = []
    for source in sources:
        for png in omr.page_images(source, pages if source.suffix.lower() == ".pdf" else None, work):
            found = omr.find_systems(png)
            log(f"{png.name}: {len(found)} line(s) of music found")
            systems.extend(_trim(crop) for crop in found)
    if not systems:
        raise SongError("No lines of music were found. Use a straight, clear scan or photo "
                        "(300 dpi or more), with the whole page in the picture.")
    return systems


def _trim(crop: Image.Image, margin: int = 12) -> Image.Image:
    gray = crop.convert("L")
    ink = gray.point(lambda v: 255 if v < 160 else 0)
    box = ink.getbbox()
    if not box:
        return gray
    l, t, r, b = box
    return gray.crop((max(0, l - margin), max(0, t - margin),
                      min(gray.width, r + margin), min(gray.height, b + margin)))


def ink_image(gray: Image.Image, color: str) -> Image.Image:
    """The printed ink in one colour on a transparent background."""
    alpha = ImageOps.invert(ImageOps.autocontrast(gray.convert("L"), cutoff=1))
    alpha = alpha.point(lambda v: 0 if v < 50 else min(255, int(v * 1.3)))  # drop paper grain
    out = Image.new("RGBA", gray.size, color)
    out.putalpha(alpha)
    return out


def stack(images: list[Image.Image]) -> Image.Image:
    width = max(im.width for im in images)
    height = sum(im.height for im in images) + GAP * (len(images) - 1)
    out = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    y = 0
    for im in images:
        out.alpha_composite(im, (0, y))
        y += im.height + GAP
    return out


def build_picture_deck(sources: list[str | Path], out: str | Path, meta: dict, theme_name: str = "light",
                       bg: str | None = None, pages: str | None = None, per_slide: int = 2,
                       verses: int = 1, log=print) -> int:
    """Make the deck; return the number of music slides."""
    check_copyright(meta, engraving=False)
    theme = make_theme(theme_name, bg)
    with tempfile.TemporaryDirectory() as tmp:
        systems = systems_from_scans([Path(s) for s in sources], pages, Path(tmp), log)
    inked = [ink_image(s, theme.music) for s in systems]
    groups = [inked[i:i + per_slide] for i in range(0, len(inked), per_slide)]

    prs = _new_deck()
    title_slide(prs, theme, meta)
    left, top, bw, bh = _music_box(theme)
    count = 0
    for verse in range(1, verses + 1):
        for group in groups:
            slide = _blank_slide(prs, theme)
            label = f"Stanza {verse}" if verses > 1 else str(meta.get("title", ""))
            _label(slide, label, theme)
            if theme.panel:
                _panel(slide)
            png = stack(group)
            w_in = min(bw, bh * png.width / png.height)
            h_in = w_in * png.height / png.width
            buf = io.BytesIO()
            png.save(buf, format="PNG")
            buf.seek(0)
            slide.shapes.add_picture(buf, Inches(left + (bw - w_in) / 2), Inches(top + (bh - h_in) / 2),
                                     Inches(w_in), Inches(h_in))
            slide.notes_slide.notes_text_frame.text = label
            count += 1
    prs.save(out)
    log(f"Wrote {out}  ({count} music slides from {len(systems)} lines of music)")
    return count

