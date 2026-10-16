from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from hymnslides.deck import build_deck, build_template
from hymnslides.songformat import SongError, load_song, song_from_dict

EXAMPLES = Path(__file__).parent.parent / "examples"
SONGS = ["blessed_assurance", "standing_on_the_promises"]
THEMES = [("light", None), ("dark", None), ("color", "#8B1A1A")]


@pytest.fixture(scope="module", params=[(s, t) for s in SONGS for t in THEMES],
                ids=lambda p: f"{p[0]}-{p[1][0]}")
def deck(request, tmp_path_factory):
    name, (theme, bg) = request.param
    song = load_song(EXAMPLES / f"{name}.yaml")
    out = tmp_path_factory.mktemp("decks") / f"{name}-{theme}.pptx"
    planned = build_deck(song, out, theme, bg, log=lambda *_: None)
    return song, planned, Presentation(str(out))


def test_every_music_slide_has_one_picture(deck):
    _, planned, prs = deck
    slides = list(prs.slides)
    assert len(slides) == len(planned) + 1
    for slide in slides[1:]:
        pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pics) == 1


def test_lyrics_at_least_30pt(deck):
    _, planned, _ = deck
    assert min(ms.lyric_pt for ms in planned) >= 30


def test_all_text_at_least_30pt(deck):
    _, _, prs = deck
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        assert run.font.size.pt >= 30, run.text


def test_refrain_after_every_stanza(deck):
    song, _, prs = deck
    labels = [s.notes_slide.notes_text_frame.text.splitlines()[0] for s in list(prs.slides)[1:]]
    sections = [l for i, l in enumerate(labels) if i == 0 or labels[i - 1] != l]
    stanzas = [s.label for s in song.structure if s.label.startswith("Stanza")]
    assert len(stanzas) >= 3
    assert sections == [x for st in stanzas for x in (st, "Refrain")]


def _tiny_song(**meta):
    return song_from_dict({
        "title": "t", "key": "C", "time": "4/4", **meta,
        "phrases": [{"name": "A", "treble": "w:C4", "bass": "w:C3"}],
        "lyrics": {"s1": {"A": "Hi"}},
        "structure": [{"label": "Stanza 1", "phrases": ["A"], "lyrics": "s1"}],
    })


def test_licensed_without_ccli_fails(tmp_path):
    song = _tiny_song(copyright_status="licensed", copyright="© 2001 Someone")
    with pytest.raises(SongError, match="ccli_song, ccli_license"):
        build_deck(song, tmp_path / "x.pptx", log=lambda *_: None)


def test_template_only_cannot_be_engraved(tmp_path):
    song = _tiny_song(copyright_status="template_only")
    with pytest.raises(SongError, match="template_only"):
        build_deck(song, tmp_path / "x.pptx", log=lambda *_: None)


def test_template_deck(tmp_path):
    out = tmp_path / "t.pptx"
    n = build_template({"title": "My Hymn", "copyright_status": "template_only"},
                       "Stanza 1:2,Refrain:2", out, "color", "#8B1A1A")
    prs = Presentation(str(out))
    assert n == 4 and len(prs.slides) == 5
    texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
    assert "Stanza 1" in texts and "Paste 2 music systems here" in texts
